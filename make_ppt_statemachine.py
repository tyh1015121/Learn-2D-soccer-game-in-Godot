# -*- coding: utf-8 -*-
"""
生成《2D 足球游戏 状态机结构解析》PPT
用法: python make_ppt_statemachine.py
输出: 状态机结构解析.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------- 配色 ----------
DARK   = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝（标题条）
ACCENT = RGBColor(0x54, 0x82, 0x35)   # 草地绿
INK    = RGBColor(0x26, 0x26, 0x26)
GRAY   = RGBColor(0x59, 0x59, 0x59)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

P_FILL, P_LINE = RGBColor(0xDA, 0xE8, 0xFC), RGBColor(0x6C, 0x8E, 0xBF)  # 玩家状态 浅蓝
P_HUB,  P_HUBL = RGBColor(0x9D, 0xC3, 0xE6), RGBColor(0x2E, 0x75, 0xB6)  # 玩家枢纽 深一档
B_FILL, B_LINE = RGBColor(0xD5, 0xE8, 0xD4), RGBColor(0x82, 0xB3, 0x66)  # 球状态 浅绿
B_HUB,  B_HUBL = RGBColor(0xA9, 0xD1, 0x8E), RGBColor(0x54, 0x82, 0x35)
G_FILL, G_LINE = RGBColor(0xE7, 0xE6, 0xE6), RGBColor(0x99, 0x99, 0x99)  # 未启用 灰
C_FILL, C_LINE = RGBColor(0xF2, 0xF7, 0xEC), RGBColor(0x82, 0xB3, 0x66)  # 容器
CODE_BG        = RGBColor(0xF5, 0xF5, 0xF0)

FONT = "Microsoft YaHei"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def style_run(run, size=12, bold=False, color=INK, font=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def set_dash(shape_or_line):
    ln = shape_or_line.line._get_or_add_ln()
    d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
    ln.append(d)


def box(slide, x, y, w, h, text, fill, line, size=11, bold=False,
        color=INK, rounded=True, dashed=False, line_w=1.2):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(line_w)
    if dashed:
        set_dash(shp)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, t in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = t
        style_run(r, size, bold, color)
    return shp


def label(slide, x, y, w, h, text, size=9, color=GRAY, bold=False,
          align=PP_ALIGN.CENTER, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    for i, t in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        style_run(r, size, bold, color, font)
    return tb


def conn(slide, x1, y1, x2, y2, color=GRAY, width=1.5, dashed=False, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if dashed:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if arrow:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    return c


def poly_arrow(slide, pts, color=GRAY, width=1.5, dashed=False):
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        conn(slide, x1, y1, x2, y2, color, width, dashed,
             arrow=(i == len(pts) - 2))


def new_slide(title=None, subtitle=None):
    slide = prs.slides.add_slide(BLANK)
    if title:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(0), Inches(0), prs.slide_width, Inches(0.75))
        bar.fill.solid()
        bar.fill.fore_color.rgb = DARK
        bar.line.fill.background()
        bar.shadow.inherit = False
        tf = bar.text_frame
        tf.margin_left = Inches(0.45)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        style_run(r, 20, True, WHITE)
        acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(0), Inches(0.75), prs.slide_width, Inches(0.06))
        acc.fill.solid()
        acc.fill.fore_color.rgb = ACCENT
        acc.line.fill.background()
        acc.shadow.inherit = False
        if subtitle:
            label(slide, 7.0, 0.24, 5.9, 0.35, subtitle, 10,
                  RGBColor(0xBD, 0xD7, 0xEE), align=PP_ALIGN.RIGHT)
    return slide


def bullets(slide, x, y, w, h, items, size=13):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        lvl, txt = item[0], item[1]
        bold = item[2] if len(item) > 2 else (lvl < 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        if lvl < 0:
            prefix = ""
        elif lvl == 0:
            prefix = "• "
        else:
            prefix = "      – "
        r = p.add_run()
        r.text = prefix + txt
        style_run(r, size - (1 if lvl > 0 else 0),
                  bold, DARK if lvl < 0 else INK)
    return tb


def code_block(slide, x, y, w, h, lines, size=11):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = CODE_BG
    shp.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.1)
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t if t else " "
        style_run(r, size, False, RGBColor(0x24, 0x29, 0x2E), MONO)
    return shp


# =====================================================================
# 1. 封面
# =====================================================================
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                        prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
bg.shadow.inherit = False
stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(4.62),
                            prs.slide_width, Inches(0.08))
stripe.fill.solid(); stripe.fill.fore_color.rgb = ACCENT; stripe.line.fill.background()
stripe.shadow.inherit = False
label(s, 1.0, 2.35, 11.33, 1.0, "2D 足球游戏 · 状态机结构解析", 40, WHITE, True)
label(s, 1.0, 3.55, 11.33, 0.5,
      "Learn-2D-soccer-game-in-Godot（Godot 4.6 · GDScript）",
      18, RGBColor(0xBD, 0xD7, 0xEE))
label(s, 1.0, 5.0, 11.33, 0.5,
      "Player / Ball 双状态机  ·  信号驱动切换  ·  工厂模式  ·  状态即节点",
      14, RGBColor(0xA9, 0xD1, 0x8E))

# =====================================================================
# 2. 整体架构：两套同构的状态机
# =====================================================================
s = new_slide("整体架构：两套同构的状态机", "scenes/characters · scenes/ball")

# Player 容器
box(s, 0.5, 1.25, 6.0, 5.0, "", RGBColor(0xEF, 0xF5, 0xFC), P_LINE, rounded=False)
label(s, 0.7, 1.35, 5.6, 0.4, "Player（CharacterBody2D）— player.gd", 13, DARK, True)
box(s, 2.15, 1.95, 2.7, 0.65, "switch_state(state, data)\n唯一切换入口", P_HUB, P_HUBL, 10.5, True)
box(s, 0.85, 3.55, 2.5, 1.0, "PlayerStateFactory\nstates: 枚举 → 状态类", P_FILL, P_LINE, 10.5)
box(s, 3.75, 3.55, 2.5, 1.0, "current_state\nPlayerState 子节点", P_FILL, P_LINE, 10.5)
conn(s, 2.85, 2.6, 2.1, 3.55)
label(s, 1.0, 2.95, 1.9, 0.5, "① get_fresh_state()\n每次 new 新实例", 8.5)
conn(s, 3.35, 4.05, 3.75, 4.05)
label(s, 3.15, 4.15, 2.0, 0.4, "② setup() + add_child", 8.5)
poly_arrow(s, [(5.0, 3.55), (5.0, 3.0), (3.5, 3.0), (3.5, 2.6)])
label(s, 4.6, 2.62, 2.0, 0.4, "③ 发信号请求切换", 8.5)
label(s, 0.85, 5.15, 5.5, 0.9,
      "状态以「子节点」形式挂在宿主身上，节点名 PlayerStateMachine:<状态>，\n"
      "通过 %AnimationPlayer / %BallDetectionArea 等唯一名访问宿主节点。",
      9, GRAY)

# Ball 容器
box(s, 6.85, 1.25, 6.0, 5.0, "", RGBColor(0xF1, 0xF8, 0xEE), B_LINE, rounded=False)
label(s, 7.05, 1.35, 5.6, 0.4, "Ball（AnimatableBody2D）— ball.gd", 13, ACCENT, True)
box(s, 8.5, 1.95, 2.7, 0.65, "switch_state(state)\n唯一切换入口", B_HUB, B_HUBL, 10.5, True)
box(s, 7.2, 3.55, 2.5, 1.0, "BallStateFactory\nstates: 枚举 → 状态类", B_FILL, B_LINE, 10.5)
box(s, 10.1, 3.55, 2.5, 1.0, "current_state\nBallState 子节点", B_FILL, B_LINE, 10.5)
conn(s, 9.2, 2.6, 8.45, 3.55)
label(s, 7.35, 2.95, 1.9, 0.5, "① get_fresh_state()\n每次 new 新实例", 8.5)
conn(s, 9.7, 4.05, 10.1, 4.05)
label(s, 9.5, 4.15, 2.0, 0.4, "② setup() + add_child", 8.5)
poly_arrow(s, [(11.35, 3.55), (11.35, 3.0), (9.85, 3.0), (9.85, 2.6)])
label(s, 10.95, 2.62, 2.0, 0.4, "③ 发信号请求切换", 8.5)
label(s, 7.2, 5.15, 5.5, 0.9,
      "结构完全一致，只有上下文不同（球多一个 carrier 引用）。\n"
      "另外球还可被外部直接切换：Player 射门/传球时调用 ball.shoot() / ball.pass_to()。",
      9, GRAY)

label(s, 0.5, 6.5, 12.3, 0.6,
      "关键：状态机不是独立框架，而是「物理体宿主 + 普通 Node 子节点」的组合——状态可以直接使用 "
      "_enter_tree / _process / _exit_tree 生命周期。",
      11, DARK, True)

# =====================================================================
# 3. 五要素拆解
# =====================================================================
s = new_slide("状态机的五个组成部分", "以 Player 侧为例，Ball 侧完全同构")
rows = [
    ("① 状态枚举", "Player.State { MOVING, TACKLING, RECOVERING, PREPPING_SHOT, SHOOTING, PASSING, HEADER, VOLLEY_KICK, BICYCLE_KICK, CHEST_CONTROL }（player.gd:13）——用枚举而不是字符串，切换时编译期可查。", P_FILL, P_LINE),
    ("② 状态基类", "PlayerState extends Node（player_state.gd）：定义 state_transition_requested 信号、setup() 注入依赖、transition_state() 发请求、on_animation_complete() 空钩子。", P_FILL, P_LINE),
    ("③ 具体状态", "每个状态一个 .gd，重写 _enter_tree()（进入时播动画/连信号）、_process()（逐帧逻辑）、on_animation_complete()（动画播完转换）。", P_FILL, P_LINE),
    ("④ 状态工厂", "PlayerStateFactory（player_state_factory.gd）：字典把枚举映射到状态类，get_fresh_state() 每次返回全新实例——切换即换实例，无残留变量。", P_FILL, P_LINE),
    ("⑤ 宿主", "Player 持有 current_state 与 state_factory；switch_state() 是所有切换的唯一入口（player.gd:45-52）。", P_HUB, P_HUBL),
]
y = 1.35
for title, desc, f, l in rows:
    box(s, 0.6, y, 2.3, 0.85, title, f, l, 13, True)
    label(s, 3.15, y + 0.06, 9.7, 0.8, desc, 11, INK, align=PP_ALIGN.LEFT)
    y += 1.13

# =====================================================================
# 4. switch_state 切换流程
# =====================================================================
s = new_slide("核心机制：switch_state() 的五步切换流程", "player.gd:45-52 / ball.gd:29-36")
steps = [
    "① 销毁旧状态：current_state.queue_free()",
    "② 工厂产出新实例：state_factory.get_fresh_state(state) —— 每次 .new()",
    "③ 注入上下文：setup(player, state_data, 动画器, 球, 检测区, 球门 …)",
    "④ 连接信号：state_transition_requested → 宿主的 switch_state",
    "⑤ 延迟挂载：call_deferred(\"add_child\", …)，命名 PlayerStateMachine:<状态>",
]
y = 1.5
for i, t in enumerate(steps):
    box(s, 0.7, y, 5.9, 0.72, t, P_FILL if i % 2 == 0 else P_HUB, P_LINE, 11.5)
    if i < 4:
        conn(s, 3.65, y + 0.72, 3.65, y + 1.02)
    y += 1.02

bullets(s, 7.0, 1.5, 5.9, 5.5, [
    (-1, "为什么这样设计？"),
    (0, "call_deferred：切换常发生在 _process / 信号回调里，此时直接改场景树不安全，延迟到帧末空闲时挂载。"),
    (0, "每次全新实例：状态内部的计时器、dribble_time 等临时变量自动归零，不存在「上次残留」。"),
    (0, "单一入口：任何转换最终都经过宿主的 switch_state()，调试时只盯一个函数。"),
    (0, "信号绑定 switch_state.bind()：信号参数（新状态、数据）原样转发。"),
    (-1, "状态死亡即回收"),
    (0, "queue_free() 后旧状态连同它 connect 的信号一起销毁，新旧状态永不共存。"),
], 12)

# =====================================================================
# 5. 信号驱动 + 数据载荷
# =====================================================================
s = new_slide("信号驱动的转换 & 跨状态传参", "player_state.gd / player_state_data.gd")

box(s, 0.7, 1.6, 3.0, 1.0, "PlayerStateMoving\n（当前状态节点）", P_FILL, P_LINE, 11)
box(s, 5.15, 1.6, 3.0, 1.0, "Player.switch_state()\n（宿主 · 唯一入口）", P_HUB, P_HUBL, 11, True)
box(s, 9.6, 1.6, 3.0, 1.0, "PlayerStatePassing\n（全新实例）", P_FILL, P_LINE, 11)
conn(s, 3.7, 2.1, 5.15, 2.1)
label(s, 3.35, 1.15, 2.3, 0.4, "emit 信号 + 目标枚举", 9)
conn(s, 8.15, 2.1, 9.6, 2.1)
label(s, 7.95, 1.15, 2.3, 0.4, "销毁旧态 / 创建新态", 9)

code_block(s, 0.7, 3.15, 6.3, 3.5, [
    "# 状态只发请求，从不自己切换（player_state.gd:25）",
    "func transition_state(new_state, data := PlayerStateData.new()):",
    "    state_transition_requested.emit(new_state, data)",
    "",
    "# 蓄力射门：用建造者把参数打包带进下一状态",
    "var data = PlayerStateData.build() \\",
    "    .set_shot_power(shot_power) \\",
    "    .set_shot_direction(shot_direction)",
    "transition_state(Player.State.SHOOTING, data)",
], 11)

bullets(s, 7.4, 3.15, 5.4, 3.6, [
    (-1, "三条规则"),
    (0, "状态之间零耦合：Moving 不知道 Passing 的存在，只认识枚举。"),
    (0, "PlayerStateData 载荷：解决「蓄力力度/方向要带到 SHOOTING」这类跨状态传参。"),
    (0, "动画驱动状态不在 _process 里切换：宿主 AnimationPlayer 播完 → on_animation_complete() 里再发转换（射门/传球/倒钩等）。"),
], 12)

# =====================================================================
# 6. 玩家状态机全景图
# =====================================================================
s = new_slide("玩家状态机：状态转换全景图", "实线 = 代码已实现 · 虚线 = 入口被注释")

box(s, 0.9, 3.1, 1.9, 0.75, "MOVING\n(默认)", P_HUB, P_HUBL, 11, True)
box(s, 3.7, 1.35, 2.0, 0.72, "PASSING", P_FILL, P_LINE, 11)
box(s, 6.4, 1.35, 2.1, 0.72, "PREPPING_SHOT", P_FILL, P_LINE, 11)
box(s, 9.3, 1.35, 1.9, 0.72, "SHOOTING", P_FILL, P_LINE, 11)

# 空中球技容器
box(s, 3.7, 3.05, 5.2, 1.85, "", C_FILL, C_LINE, rounded=False)
label(s, 3.8, 3.1, 5.0, 0.3, "空中球技（球在空中且可交互时按射门）", 9, ACCENT, True)
box(s, 3.95, 3.5, 1.5, 0.6, "HEADER", B_FILL, B_LINE, 10)
box(s, 5.6, 3.5, 1.6, 0.6, "VOLLEY_KICK", B_FILL, B_LINE, 10)
box(s, 7.35, 3.5, 1.5, 0.6, "BICYCLE_KICK", B_FILL, B_LINE, 10)
label(s, 3.9, 4.25, 4.9, 0.55,
      "移动中 → 头球 ；静止面向球门 → 凌空 ；静止背向 → 倒钩\n（命中加成 ×1.3 / ×1.5 / ×2.0）",
      8.5, GRAY)

box(s, 4.6, 5.5, 1.9, 0.7, "TACKLING", G_FILL, G_LINE, 11, dashed=True)
box(s, 9.3, 5.5, 1.9, 0.7, "RECOVERING", P_FILL, P_LINE, 11)

conn(s, 1.55, 3.1, 3.7, 1.75)                                  # MOVING→PASSING
label(s, 2.05, 1.95, 1.7, 0.35, "有球·按传球键", 8.5)
conn(s, 3.7, 2.0, 2.2, 3.1)                                    # PASSING→MOVING
label(s, 2.55, 2.55, 1.5, 0.35, "动画完成·传球", 8.5)
conn(s, 2.35, 3.1, 7.1, 2.07)                                  # MOVING→PREPPING
label(s, 4.35, 2.3, 2.2, 0.35, "有球·按下射门键", 8.5)
conn(s, 8.5, 1.71, 9.3, 1.71)                                  # PREPPING→SHOOTING
label(s, 7.7, 1.82, 2.2, 0.35, "松开射门+PlayerStateData", 8)
poly_arrow(s, [(10.25, 1.35), (10.25, 1.12), (1.35, 1.12), (1.35, 3.1)])  # SHOOTING→MOVING
label(s, 5.0, 0.9, 3.4, 0.22, "人类玩家 · 动画完成", 8.5)
conn(s, 10.25, 2.07, 10.25, 5.5)                               # SHOOTING→RECOVERING
label(s, 10.4, 3.5, 1.4, 0.35, "CPU·动画完成", 8.5, align=PP_ALIGN.LEFT)
conn(s, 2.8, 3.42, 3.7, 3.55)                                  # MOVING→空中技
label(s, 2.45, 3.98, 1.6, 0.35, "空中球·按射门", 8)
conn(s, 8.9, 4.4, 9.55, 5.5)                                   # 空中技→RECOVERING
label(s, 8.35, 4.85, 1.8, 0.35, "动画完成/落地", 8.5)
poly_arrow(s, [(1.85, 3.85), (1.85, 5.85), (4.6, 5.85)], dashed=True)     # MOVING→TACKLING
label(s, 1.95, 5.42, 2.6, 0.35, "铲球（入口在 moving.gd 被注释）", 8)
conn(s, 6.5, 5.85, 9.3, 5.85)                                  # TACKLING→RECOVERING
label(s, 6.95, 5.44, 2.0, 0.35, "铲球停稳 +200ms", 8.5)
poly_arrow(s, [(10.25, 6.2), (10.25, 6.9), (1.2, 6.9), (1.2, 3.85)])      # RECOVERING→MOVING
label(s, 4.9, 6.62, 2.6, 0.25, "500ms 硬直结束", 8.5)

label(s, 0.6, 7.08, 12.2, 0.35,
      "注：CHEST_CONTROL 已注册但为空实现，没有任何状态会切换到它；TACKLING 的触发代码在 player_state_moving.gd:33-34 被注释掉。",
      8.5, GRAY, align=PP_ALIGN.LEFT)

# =====================================================================
# 7. 玩家各状态职责
# =====================================================================
s = new_slide("玩家状态机：各状态职责", "scenes/characters/character_states/")
bullets(s, 0.6, 1.2, 12.2, 6.0, [
    (-1, "MOVING（默认，player_state_moving.gd）"),
    (0, "每帧经 KeyUtils 读输入 → velocity；带球时按「传球/射门」切走；球在空中可交互时按射门触发空中技；CPU 分支目前是空壳。"),
    (-1, "PASSING（player_state_passing.gd）"),
    (0, "进入即播 kick 动画并停下；动画结束时在 TeammateDetectionArea 视野内找最近队友，ball.pass_to() 传给他，否则沿朝向直塞。"),
    (-1, "PREPPING_SHOT → SHOOTING（蓄力射门两段式）"),
    (0, "按住射门蓄力、方向键微调方向；松开时按 ease 曲线算力度加成，打包 PlayerStateData 进入 SHOOTING；SHOOTING 播 kick，动画完成瞬间 ball.shoot()。"),
    (-1, "TACKLING / RECOVERING（铲球与硬直）"),
    (0, "TACKLING：地面摩擦 250 减速到停，停稳 200ms 后转 RECOVERING；RECOVERING：清零速度硬直 500ms 回 MOVING。"),
    (-1, "HEADER / VOLLEY_KICK / BICYCLE_KICK（空中技）"),
    (0, "进入时连 BallDetectionArea 的 body_entered，球进入且高度落在各自窗口内（10-30 / 1-20 / 5-25）就 contact_ball.shoot()，力度加成 ×1.3 / ×1.5 / ×2.0；动画完成或落地后转 RECOVERING。"),
], 12)

# =====================================================================
# 8. 球状态机全景图
# =====================================================================
s = new_slide("球状态机：状态转换全景图", "scenes/ball/ —— 只有 3 个状态")

box(s, 1.6, 3.3, 2.2, 0.8, "CARRIED\n(被带球)", B_FILL, B_LINE, 11, True)
box(s, 6.0, 1.5, 2.2, 0.8, "SHOT\n(射门飞行)", B_FILL, B_LINE, 11, True)
box(s, 10.4, 3.3, 2.2, 0.8, "FREEFORM\n(自由·默认)", B_HUB, B_HUBL, 11, True)

conn(s, 3.3, 3.3, 6.3, 2.3)                                    # CARRIED→SHOT
label(s, 3.6, 2.3, 2.6, 0.35, "射门 ball.shoot()", 9)
conn(s, 7.5, 2.3, 11.0, 3.3)                                   # SHOT→FREEFORM
label(s, 8.5, 2.3, 2.8, 0.35, "1 秒超时 / 撞墙提前反弹", 9)
conn(s, 10.4, 3.62, 3.8, 3.62)                                 # FREEFORM→CARRIED
label(s, 5.2, 3.18, 3.6, 0.35, "玩家进入 PlayerDetectionArea（信号）", 9)
conn(s, 3.8, 3.98, 10.4, 3.98)                                 # CARRIED→FREEFORM
label(s, 5.2, 4.1, 3.6, 0.35, "传球 ball.pass_to()", 9)

label(s, 1.35, 4.45, 2.7, 1.0,
      "贴住 carrier 脚下，\ncos 函数模拟运球抖动", 9, GRAY)
label(s, 5.75, 2.55, 2.7, 0.8,
      "固定高度飞行，\n期间不可被拾取", 9, GRAY)
label(s, 10.15, 4.45, 2.7, 1.0,
      "摩擦减速 + 重力弹跳；\n撞墙反弹并自转换回 FREEFORM", 9, GRAY)

box(s, 1.2, 5.7, 11.0, 1.15,
    "注意两条切换路径：① 状态内部发信号（FREEFORM→CARRIED、SHOT→FREEFORM）；"
    "② 宿主直接调用 —— 玩家射门/传球时由 ball.shoot() / ball.pass_to() 里直接 switch_state()，"
    "以及反弹时 move_and_bounce() 内部自切。",
    RGBColor(0xFF, 0xF2, 0xCC), RGBColor(0xD6, 0xB6, 0x56), 11)

# =====================================================================
# 9. 球状态职责 + 假高度
# =====================================================================
s = new_slide("球状态机：职责与「假高度」物理", "ball_state.gd / ball_state_*.gd")
bullets(s, 0.6, 1.2, 12.2, 3.6, [
    (-1, "FREEFORM（ball_state_freeform.gd）"),
    (0, "按所在高度选地面/空气摩擦减速；process_gravity 带 BOUNCINESS=0.8 弹跳；move_and_bounce 撞墙反弹；玩家进入检测区 → CARRIED；can_air_interact() 返回 true（只有此时能触发空中技）。"),
    (-1, "CARRIED（ball_state_carried.gd）"),
    (0, "每帧把球钉在 carrier 脚下偏移处，cos(dribble_time×10)×3 左右抖动模拟运球；自己不切状态，等玩家射门/传球时被宿主切走。"),
    (-1, "SHOT（ball_state_shot.gd）"),
    (0, "固定高度 5 飞行 1 秒，精灵压扁 0.8 倍表现速度感；超时转 FREEFORM；_exit_tree 恢复缩放。"),
], 12)
code_block(s, 0.6, 4.9, 6.4, 1.9, [
    "# 2D 游戏里的「假高度」：height 只是数值，不参与碰撞",
    "ball.height_velocity -= GRAVITY * delta   # 重力",
    "ball.height += ball.height_velocity       # 积分",
    "if ball.height < 0:                       # 落地反弹",
    "    ball.height_velocity = -v * bounciness",
    "ball_sprite.position = Vector2.UP * height  # 仅视觉上移",
], 10.5)
bullets(s, 7.4, 4.9, 5.4, 2.0, [
    (-1, "为什么要假高度？"),
    (0, "俯视角 2D 没有真正的 Z 轴；用 height 数值 + 精灵偏移模拟抛物线，空中技再用高度窗口（如头球 10~30）判断「够得着吗」。"),
], 11.5)

# =====================================================================
# 10. 新增状态 & 总结
# =====================================================================
s = new_slide("如何新增一个状态 & 设计要点总结")
box(s, 0.6, 1.25, 5.9, 0.6, "新增状态的固定三步", P_HUB, P_HUBL, 13, True)
bullets(s, 0.6, 2.05, 6.1, 2.4, [
    (0, "新建 player_state_xxx.gd，extends PlayerState，重写需要的生命周期函数。"),
    (0, "在 Player.State 枚举里加一项。"),
    (0, "在 PlayerStateFactory.states 字典里注册：枚举 → 类。"),
    (0, "Ball 侧同理（Ball.State + BallStateFactory）。"),
], 12)
code_block(s, 0.6, 4.7, 6.1, 2.1, [
    "# 工厂模式让注册即接入",
    "states = {",
    "    Player.State.MOVING: PlayerStateMoving,",
    "    Player.State.XXX:    PlayerStateXxx,   # ← 加一行",
    "}",
], 10.5)
box(s, 7.1, 1.25, 5.7, 0.6, "这套状态机值得借鉴的点", B_HUB, B_HUBL, 13, True)
bullets(s, 7.1, 2.05, 5.8, 4.8, [
    (0, "状态即节点：免生命周期框架，_enter_tree / _process / _exit_tree 天然就是 enter / update / exit。"),
    (0, "信号驱动、单一入口：状态零耦合，切换全过宿主，好调试。"),
    (0, "工厂 + 枚举：每次全新实例，无残留；新增状态三步搞定。"),
    (0, "载荷对象 PlayerStateData：跨状态传参不显式引用对方。"),
    (0, "动画驱动转换：表现（动画播完）和逻辑（何时切换）自动同步。"),
    (0, "代价：状态多了以后转换关系只存在于代码里——这张全景图就是文档。"),
], 12)

prs.save("状态机结构解析.pptx")
print("OK: 状态机结构解析.pptx  slides =", len(prs.slides._sldIdLst))
